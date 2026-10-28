# /// script
# requires-python = ">=3.11"
# dependencies = ["typer>=0.12", "pydantic>=2.6", "markdown-it-py>=3", "pyyaml>=6", "python-docx>=1.2", "openpyxl>=3.1"]
# ///
"""The run's changes as a workbook the authors filter instead of reading one comment per edit.

    uv run --script export_changes.py RUN/manuscript-dashboard/findings.json \
        --revision RUN/manuscript-improve/revision --output RUN/manuscript-dashboard/manuscript-changes.xlsx

Under the default `edits: apply` the run applies every change that needs no decision of the authors, and the
Word review (`export_docx.py --detail decisions`) carries only what does; this workbook is where the applied
changes are listed. It opens in Excel, LibreOffice and Google Sheets, one sheet per question:

- **Applied**: every edit in this version, with the stage, section, the text before and after, why, and the
  problems it answers; an edit whose original wording a later change restored says so.
- **For you**: what needs the authors, the same items the decisions document comments on: each decision only
  they can make, each question, and each proposal in this version or waiting for their answer.
- **Open problems**: every problem no applied edit answers, with its severity and when to fix it, whatever its
  severity (the Word review keeps only the blockers and major ones).
- **Held**: the proposals left out of this version and why (another file, an alternative chosen, a decision).
- **Co-authors' comments**, only when the run tracked them (manuscript-coauthors): each comment or suggested edit
  on the previous release, with its author, status, the passage before and now, why, and what could still be done.

Each header row stays in view and every column filters. An existing output is never overwritten: the message
names the next attempt (manuscript-changes.2.xlsx). `runs.py handover` puts the newest one beside the
manuscript as STEM.changes.xlsx.
"""
from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path
from typing import Annotated

import typer
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve()
sys.path.insert(0, str(HERE.parent))
import dashboard as viz  # noqa: E402
import export_docx as words  # noqa: E402

# Column widths in characters; long prose wraps inside them.
WIDTHS = {"Id": 14, "Stage": 18, "Section": 28, "Change": 40, "Was": 60, "Now": 60, "Why": 60, "Problems": 40,
          "Status": 22, "Kind": 22, "What": 60, "Question or reason": 60, "Severity": 18, "When": 26,
          "Problem": 50, "Detail": 60, "Fix": 50, "Found by": 26, "Answered by": 18, "Author": 22, "Comment": 60,
          "What could be done": 50}
# A co-author comment's status and kind in the words the dashboard uses.
COAUTHOR_STATUS = {"addressed": "Addressed", "partly": "Partly addressed", "open": "Open", "declined": "Declined",
                   "not-on-point": "Not on point"}
COAUTHOR_KIND = {"comment": "Comment", "suggestion": "Suggested edit"}
HEADER_FILL = PatternFill("solid", fgColor="DDE4EE")
app = typer.Typer(add_completion=False)


def cell(text: object) -> object:
    """A value Excel accepts: control characters dropped, None as an empty cell."""
    if text is None:
        return ""
    return ILLEGAL_CHARACTERS_RE.sub("", text) if isinstance(text, str) else text


def sheet(book: Workbook, title: str, header: list[str], rows: list[list[object]]) -> None:
    ws = book.create_sheet(title)
    ws.append(header)
    for row in rows:
        ws.append([cell(value) for value in row])
    for column, name in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(column)].width = WIDTHS.get(name, 20)
        head = ws.cell(row=1, column=column)
        head.font, head.fill = Font(bold=True), HEADER_FILL
    for line in ws.iter_rows(min_row=2):
        for value in line:
            value.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def coauthor_rows(data: viz.Inventory) -> list[list[object]]:
    """One row per co-author comment: what was asked, its status, the passage before and now, why and what next. A
    suggestion's edits follow its comment, so a filtered sheet still says what it proposed."""
    rows = []
    for item in data.coauthors.items if data.coauthors else []:
        edits = [" ".join(part for part in (f"Delete: \u201c{e.deleted}\u201d" if e.deleted else "",
                                            f"Insert: \u201c{e.inserted}\u201d" if e.inserted else "") if part)
                 for e in item.edits]
        asked = "\n".join(text for text in (item.comment, *edits) if text)
        rows.append([item.id, item.author, COAUTHOR_KIND[item.kind], item.where, asked, COAUTHOR_STATUS[item.status],
                     item.before, item.now or item.after, item.why, item.next])
    return rows


def build(data: viz.Inventory, revision: dict | None) -> tuple[Workbook, dict]:
    changes = {c.id: c for c in data.changes}
    findings = {f.id: f for f in data.findings}

    def problems(change: viz.Change) -> str:
        return "; ".join(f"{findings[fid].summary.rstrip('.')} ({fid})" if fid in findings else fid for fid in change.finding_ids)

    if revision is not None:
        # The corrected version's applied edits, then those applied to the run's other files (a supplement).
        applied_ids = list(revision.get("applied", [])) + [item["id"] for item in revision.get("other_files", [])
                                                           if item["status"] == "applied"]
        in_text = set(applied_ids) | set(revision.get("proposed_in_this_version", []))
        restored = {cid: item["why"] for item in revision.get("restored", []) for cid in item["ids"]}
    else:
        applied_ids = [c.id for c in data.changes if c.status == "applied"]
        in_text, restored = set(applied_ids), {}
    applied = [[cid, words.stage_of(cid), changes[cid].section, changes[cid].summary,
                "original wording restored" if cid in restored else "applied",
                changes[cid].before, changes[cid].after, changes[cid].why or restored.get(cid, ""), problems(changes[cid])]
               for cid in [*applied_ids, *(r for r in restored if r not in applied_ids)] if cid in changes]

    asked: dict[str, list[str]] = {}
    for item in (revision or {}).get("comments", []):
        for cid in item["ids"]:
            asked.setdefault(cid, []).append(item["question"])
    for change in data.changes:  # without a corrected version, a change's own question still asks
        if change.comment and change.id not in asked:
            asked[change.id] = [change.comment]
    for_you: list[list[object]] = []
    answered: set[str] = {fid for cid in applied_ids if cid in changes for fid in changes[cid].finding_ids}
    for decision in viz.decision_list(data):
        answered |= set(decision["finding_ids"])
        for_you.append(["decision", decision["id"], ", ".join(dict.fromkeys(findings[f].section for f in decision["finding_ids"]
                                                                            if f in findings)),
                        decision["summary"], "", "", "A decision only you can make.",
                        ", ".join(decision["finding_ids"])])
    held = {item["id"]: item["why"] for item in (revision or {}).get("held", [])}
    for change in data.changes:
        questions = asked.get(change.id, [])
        if change.status == "applied" and not questions:
            continue
        if change.status == "applied":
            kind = "question on an applied edit"
        elif change.id in held:
            continue  # on the Held sheet, with its reason
        elif change.id in in_text:
            kind = "proposal in this version"
        else:
            kind = "proposal waiting for your answer" if questions else "proposal not in this version"
        answered |= set(change.finding_ids)
        for_you.append([kind, change.id, change.section, change.summary, change.before, change.after,
                        " ".join([*(f"Question: {q}" for q in questions), change.why]).strip(), ", ".join(change.finding_ids)])

    open_rows = [[f.id, words.SEVERITY[f.severity], words.WHEN[f.stage], f.section, f.summary, f.detail, f.action,
                  ", ".join(words.skill_name(s) for s in f.skills),
                  next((cid for cid, c in changes.items() if f.id in c.finding_ids and cid not in applied_ids), "")]
                 for f in data.findings
                 if f.assessment != "dismissed" and f.resolution != "addressed"
                 and f.id not in {fid for cid in applied_ids if cid in changes for fid in changes[cid].finding_ids}]
    alternatives = {item["id"]: f"An alternative to {item['chosen']}, which this version uses."
                    for item in (revision or {}).get("alternatives", [])}
    not_placed = {item["id"]: item["why"] for item in (revision or {}).get("not_placed", [])}
    held_rows = [[cid, words.stage_of(cid), changes[cid].section, changes[cid].summary, why,
                  changes[cid].before, changes[cid].after]
                 for cid, why in {**held, **alternatives, **not_placed}.items() if cid in changes]

    book = Workbook()
    book.remove(book.active)
    sheet(book, "Applied", ["Id", "Stage", "Section", "Change", "Status", "Was", "Now", "Why", "Problems"], applied)
    sheet(book, "For you", ["Kind", "Id", "Section", "What", "Was", "Now", "Question or reason", "Problems"], for_you)
    sheet(book, "Open problems", ["Id", "Severity", "When", "Section", "Problem", "Detail", "Fix", "Found by", "Answered by"],
          open_rows)
    sheet(book, "Held", ["Id", "Stage", "Section", "Change", "Question or reason", "Was", "Now"], held_rows)
    coauthors = coauthor_rows(data)
    if coauthors:
        sheet(book, "Co-authors' comments", ["Id", "Author", "Kind", "Section", "Comment", "Status", "Was", "Now", "Why",
                                             "What could be done"], coauthors)
    book.properties.title = data.title
    book.properties.creator = "manuscript-improve"
    return book, {"applied": sum(r[4] == "applied" for r in applied), "restored": sum(r[4] != "applied" for r in applied),
                  "for_you": len(for_you), "open_problems": len(open_rows), "held": len(held_rows),
                  "coauthors": len(coauthors)}


@app.command()
def main(inventory: Annotated[Path, typer.Argument(exists=True, dir_okay=False)],
         output: Annotated[Path, typer.Option("--output", help="The .xlsx to write")],
         revision: Annotated[Path | None, typer.Option(help="revise.py's output folder, for the edits in this version",
                                                       exists=True, file_okay=False)] = None) -> None:
    if output.exists():
        raise typer.BadParameter(f"{output.name} exists; write the next attempt as {viz.next_attempt(output, folder=False).name}.")
    if output.suffix.lower() != ".xlsx":
        raise typer.BadParameter("The output is an .xlsx file.")
    try:
        data = viz.Inventory.model_validate_json(inventory.read_text(encoding="utf-8"))
        corrected = json.loads((revision / "revision.json").read_text(encoding="utf-8")) if revision else None
        if corrected and corrected.get("inventory_sha256") != hashlib.sha256(inventory.read_bytes()).hexdigest():
            raise ValueError("The corrected version was built from another inventory; rebuild it with revise.py from this one.")
        book, counts = build(data, corrected)
    except (ValueError, OSError) as exc:
        raise typer.BadParameter(str(exc)) from exc
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(str(output))
    typer.echo(json.dumps({"output": str(output.resolve()), **counts}, indent=2))


if __name__ == "__main__":
    for stream in (sys.stdout, sys.stderr):  # UTF-8 on every platform, a Windows console included
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    app()
