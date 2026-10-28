# /// script
# requires-python = ">=3.11"
# dependencies = ["typer>=0.12", "pydantic>=2.6", "markdown-it-py>=3", "python-docx>=1.2"]
# ///
"""The run as a Word document to read and share: the corrected version, and what needs the authors as comments
in the margin.

    uv run --script export_docx.py RUN/manuscript-dashboard/findings.json \
        --revision RUN/manuscript-improve/revision --output RUN/manuscript-dashboard/manuscript-review.docx

`--detail decisions` (the default) is the document for the authors' attention: with --revision the text is the
corrected version, an applied edit that asks nothing is already in it, and only the proposals and the applied
edits that carry a question are tracked changes. The comments: each question for the authors on its edit, each
proposal's reasoning, each decision only the authors can make (on the passage of the first problem it covers),
and each open blocker or major problem that no applied edit, proposal or decision already answers. The applied
edits and every other problem are in the changes workbook (export_changes.py) and the dashboard, not here.

`--detail everything` is the whole review, for authors who approve every change (`edits: propose`): every
applied edit and chosen proposal is a tracked change signed by the stage that made it, each problem a comment on
the passage it quotes (or on its section's heading when the quote is not in the text), and each edit's reasoning
a comment. Without --revision, the text is the reviewed version.

Each comment opens with what it says, and the problem's or edit's ID follows in parentheses. Nothing is added to
the text itself: a question is a comment, never a placeholder, and an image that could not be embedded is named
in a comment where it stands.

An existing output is never overwritten: the message names the next attempt (manuscript-review.2.docx).

Google Docs keeps both: upload the file to Google Drive and open it with Google Docs, and the comments arrive
as comments and the tracked changes as suggestions.
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Annotated

import typer
from docx import Document as new_document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve()
sys.path.insert(0, str(HERE.parent))
import dashboard as viz  # noqa: E402

sys.path.insert(0, str(HERE.parents[2] / "manuscript-improve" / "scripts"))
import revise  # noqa: E402

SENTINEL = re.compile("(\\d+)|(\\d+)|[]")
SENTENCE = re.compile(r"(?<=[.!?])(?=\s)")
PREFIX = {"IC": "integrity-check", "RF": "reference-check", "RP": "reporting-check", "PR": "peer-review",
          "PF": "policy-fit", "JS": "journal-select", "JA": "journal-adapt", "PE": "prose-edit", "DS": "deslop",
          "ST": "story", "FG": "figures", "IM": "improve", "RR": "reviewer-response", "RM": "revision-marks", "MP": "mark-people"}
NAME = {"integrity-check": "Integrity check", "reference-check": "Reference check", "reporting-check": "Reporting check",
        "peer-review": "Simulated reviewers", "policy-fit": "Policy fit", "journal-select": "Journal selection",
        "journal-adapt": "Journal adaptation", "prose-edit": "Prose edit", "deslop": "Deslop", "story": "Story",
        "figures": "Figures and tables", "improve": "Workflow", "reviewer-response": "Reviewer response",
        "revision-marks": "Revision marks", "mark-people": "People behind the marks"}
# The words the authors read for each severity and stage, as in the dashboard (stage-findings.md).
SEVERITY = {"blocker": "Blocker", "major": "Major", "minor": "Minor", "info": "Optional", "unknown": "Could not be checked"}
WHEN = {"now": "fix now", "writing": "then the writing", "journal": "once you have chosen the journal",
        "final": "before submission, once the code is final", "unknown": "stage not set"}
IMAGES = {".png", ".jpg", ".jpeg", ".gif", ".bmp"}
# The problems the decisions document still comments on when no edit, proposal or decision answers them.
NEEDS_ATTENTION = {"blocker", "major"}
DETAIL = ("decisions", "everything")
app = typer.Typer(add_completion=False)


def stage_of(change_id: str) -> str:
    """The stage that made an edit, from its id: "ST-C3" from its prefix, "prose-edit.C1" from its skill."""
    skill = re.match(r"([a-z-]+)\.C\d+$", change_id)
    return NAME.get(skill.group(1) if skill else PREFIX.get(change_id.split("-")[0], ""), "Workflow")


def skill_name(skill: str) -> str:
    return NAME.get(skill.removeprefix("manuscript-"), skill)


def plural(n: int, word: str) -> str:
    return f"{n} {word}{'' if n == 1 else 's'}"


def initials(name: str) -> str:
    return "".join(word[0] for word in name.split()[:2]).upper()


def normalized(text: str) -> tuple[str, list[int]]:
    """Text as quotes are compared: no Markdown marks, one space for any whitespace, lower case; with each
    character's position in the original."""
    out, where = [], []
    for at, char in enumerate(text):
        if char in "*_`#>\\^~":
            continue
        if char.isspace():
            if out and out[-1] == " ":
                continue
            char = " "
        out.append(char.lower())
        where.append(at)
    return "".join(out), where


class Builder:
    """Markdown with the change sentinels of revise.marked() as a Word document. Every run is recorded with its kind
    (equal, del or ins) and edit, so that comments can anchor on runs and edits become tracked changes last."""

    def __init__(self, folders: list[Path]) -> None:
        self.doc = new_document()
        self.doc.styles["Normal"].font.size = Pt(11)
        self.folders = folders
        self.runs: list[tuple[object, str, int | None, str]] = []  # run, kind, edit number, text
        self.headings: list[tuple[str, object]] = []
        self.change: tuple[str, int] | None = None
        self.change_id = 0
        self.missing_images: list[tuple[object, str]] = []  # an empty run where an image stands, and its source

    # ----- text -----
    def text(self, paragraph, text: str, fmt: dict, link=None) -> None:
        at = 0
        for match in SENTINEL.finditer(text):
            self.plain(paragraph, text[at:match.start()], fmt, link)
            at = match.end()
            if match.group(1) is not None:
                self.change = ("del", int(match.group(1)))
            elif match.group(2) is not None:
                self.change = ("ins", int(match.group(2)))
            else:
                self.change = None
        self.plain(paragraph, text[at:], fmt, link)

    def plain(self, paragraph, text: str, fmt: dict, link) -> None:
        for part in SENTENCE.split(text) if not fmt.get("code") else [text]:
            if not part:
                continue
            run = paragraph.add_run(part)
            run.bold, run.italic = fmt.get("bold") or None, fmt.get("italic") or None
            # Both share one vertical-alignment element: setting the other to None would clear it.
            if fmt.get("sup"):
                run.font.superscript = True
            elif fmt.get("sub"):
                run.font.subscript = True
            if fmt.get("code"):
                run.font.name = "Consolas"
            if link is not None:
                run.font.color.rgb, run.font.underline = RGBColor(0x05, 0x63, 0xC1), True
                link.append(run._r)
            kind, number = self.change or ("equal", None)
            self.runs.append((run, kind, number, part))

    def hyperlink(self, paragraph, url: str):
        element = OxmlElement("w:hyperlink")
        element.set(qn("r:id"), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
        paragraph._p.append(element)
        return element

    def image(self, paragraph, token) -> None:
        source = token.attrGet("src") or ""
        name = source.split("#")[0].split("?")[0]
        for folder in self.folders:
            path = (folder / name).resolve()
            if not re.match(r"(?i)^[a-z][a-z0-9+.-]*:", source) and path.is_file() and path.suffix.lower() in IMAGES:
                picture = paragraph.add_run().add_picture(str(path), width=Inches(6))
                if token.content:  # the alt text, from the image's brackets or its pandoc attribute block
                    picture._inline.docPr.set("descr", token.content)
                return
        # Nothing is added to the text: an empty run marks the place, and a comment on it names the image.
        self.missing_images.append((paragraph.add_run(""), source))

    def inline(self, paragraph, children) -> None:
        fmt, link = {}, None
        for token in children or []:
            kind = token.type
            if kind == "text":
                self.text(paragraph, token.content, fmt, link)
            elif kind == "code_inline":
                self.text(paragraph, token.content, {**fmt, "code": True}, link)
            elif kind == "softbreak":
                self.text(paragraph, " ", fmt, link)
            elif kind == "hardbreak":
                paragraph.add_run().add_break()
            elif kind in ("strong_open", "strong_close"):
                fmt["bold"] = kind == "strong_open"
            elif kind in ("em_open", "em_close"):
                fmt["italic"] = kind == "em_open"
            elif kind in ("sup_open", "sup_close", "sub_open", "sub_close"):
                fmt[kind[:3]] = kind.endswith("_open")
            elif kind == "link_open":
                href = token.attrGet("href") or ""
                link = self.hyperlink(paragraph, href) if href and not href.startswith("#") else None
            elif kind == "link_close":
                link = None
            elif kind == "image":
                self.image(paragraph, token)

    # ----- blocks -----
    def render(self, markdown: str) -> None:
        tokens = viz.markdown_parser().parse(markdown)
        styles = {style.name for style in self.doc.styles}
        lists: list[str] = []
        quote = 0
        i = 0
        while i < len(tokens):
            token = tokens[i]
            kind = token.type
            if kind == "heading_open":
                level = int(token.tag[1])
                paragraph = self.doc.add_paragraph(style="Title" if level == 1 else f"Heading {min(level - 1, 9)}")
                start = len(self.runs)
                self.inline(paragraph, tokens[i + 1].children)
                self.headings.append((SENTINEL.sub("", tokens[i + 1].content), self.runs[start][0] if len(self.runs) > start else None))
                i += 3
                continue
            if kind in ("bullet_list_open", "ordered_list_open"):
                lists.append("List Bullet" if kind == "bullet_list_open" else "List Number")
            elif kind in ("bullet_list_close", "ordered_list_close"):
                lists.pop()
            elif kind == "blockquote_open":
                quote += 1
            elif kind == "blockquote_close":
                quote -= 1
            elif kind == "paragraph_open":
                style = (lists[-1] + (f" {len(lists)}" if len(lists) > 1 else "")) if lists else ("Quote" if quote else None)
                paragraph = self.doc.add_paragraph(style=style if style in styles else None)
                start = len(self.runs)
                self.inline(paragraph, tokens[i + 1].children)
                runs = self.runs[start:]
                # A short paragraph in bold ("**Code Availability**") is a section heading in many manuscripts.
                if runs and not lists and all(run.bold for run, *_ in runs) and len("".join(t for *_, t in runs)) <= 80:
                    self.headings.append((SENTINEL.sub("", tokens[i + 1].content).strip("*_ :"), runs[0][0]))
                i += 3
                continue
            elif kind in ("fence", "code_block"):
                paragraph = self.doc.add_paragraph()
                for number, line in enumerate(token.content.rstrip("\n").split("\n")):
                    if number:
                        paragraph.add_run().add_break()
                    self.text(paragraph, line, {"code": True})
            elif kind == "table_open":
                i = self.table(tokens, i)
                continue
            elif kind == "hr":
                self.doc.add_paragraph()
            i += 1

    def table(self, tokens, i: int) -> int:
        rows, row = [], None
        while tokens[i].type != "table_close":
            kind = tokens[i].type
            if kind == "tr_open":
                row = []
            elif kind == "tr_close":
                rows.append(row)
            elif kind in ("th_open", "td_open"):
                row.append((kind == "th_open", tokens[i + 1].children))
                i += 2
            i += 1
        width = max(len(r) for r in rows) if rows else 0
        if width:
            grid = self.doc.add_table(rows=len(rows), cols=width)
            grid.style = "Table Grid"
            for r, cells in enumerate(rows):
                for c, (head, children) in enumerate(cells):
                    paragraph = grid.cell(r, c).paragraphs[0]
                    start = len(self.runs)
                    self.inline(paragraph, children)
                    if head:
                        for run, *_ in self.runs[start:]:
                            run.bold = True
        return i + 1

    # ----- comments and tracked changes -----
    def locate(self, quote: str) -> list | None:
        """The runs of the reviewed text (unchanged and deleted words) that hold the quote."""
        owners, stream = [], []
        for index, (run, kind, _, text) in enumerate(self.runs):
            if kind != "ins":
                stream.append(text)
                owners.extend([index] * len(text))
        haystack, where = normalized("".join(stream))
        for needle in (normalized(quote)[0].strip(), normalized(quote[:120])[0].strip()):
            if len(needle) >= 12 and (at := haystack.find(needle)) >= 0:
                first, last = owners[where[at]], owners[where[at + len(needle) - 1]]
                return [self.runs[first][0], self.runs[last][0]]
        return None

    def heading_for(self, section: str):
        number = re.match(r"\s*(?:section\s+)?(\d+(?:\.\d+)*)", section, re.I)
        for text, run in self.headings:
            if run is None:
                continue
            if number and re.match(rf"\s*{re.escape(number.group(1))}\b", text):
                return run
        lower = section.lower()
        for text, run in self.headings:
            name = re.sub(r"^\d+(?:\.\d+)*\.?\s*", "", text).lower()
            if run is not None and len(name) > 3 and (name in lower or lower in name):
                return run
        return next((run for _, run in self.headings if run is not None), self.runs[0][0] if self.runs else None)

    def comment(self, runs, text: str, author: str) -> None:
        self.doc.add_comment(runs, text=text, author=author, initials=initials(author))

    def track(self, author_of: dict[int, str], date: str) -> int:
        """Wrap every inserted and deleted run as a tracked change of the edit's stage."""
        count = 0
        for run, kind, number, _ in self.runs:
            if kind == "equal":
                continue
            element = OxmlElement("w:ins" if kind == "ins" else "w:del")
            self.change_id += 1
            element.set(qn("w:id"), str(self.change_id))
            element.set(qn("w:author"), author_of.get(number, "Workflow"))
            element.set(qn("w:date"), date)
            r = run._r
            r.addprevious(element)
            element.append(r)
            if kind == "del":
                for t in r.findall(qn("w:t")):
                    t.tag = qn("w:delText")
            count += 1
        return count


def build(data: viz.Inventory, base: Path, revision: dict | None, problems: bool, reasons: bool,
          detail: str = "everything") -> tuple[Builder, dict]:
    if detail not in DETAIL:
        raise ValueError(f"--detail is one of {', '.join(DETAIL)}, not {detail!r}.")
    decisions_only = detail == "decisions"
    # Questions by the edit's IDs: the edit numbers change when accepted edits leave the tracked changes.
    asked_for: dict[tuple[str, ...], list[str]] = {tuple(e["ids"]): e.get("comments", []) for e in (revision or {}).get("edits", [])}
    asking = {cid for item in (revision or {}).get("comments", []) for cid in item["ids"]}
    accepted: list[str] = []
    if revision is not None:
        if "pieces" not in revision:
            raise ValueError("This corrected version was built before exports were possible: rebuild it with revise.py.")
        document = revise.Document("")
        pieces = []
        for p in revision["pieces"]:
            if "text" in p:
                pieces.append(revise.Piece(text=p["text"]))
            elif decisions_only and p["status"] == "applied" and not asking & set(p["ids"]):
                # Applied and asking nothing: already in the text, and listed in the changes workbook.
                pieces.append(revise.Piece(text=p["new"]))
                accepted.extend(p["ids"])
            else:
                pieces.append(revise.Piece(old=p["old"], new=p["new"], base=p["new"], ids=p["ids"], status=p["status"]))
        document.pieces = pieces
        markdown, edits = revise.marked(document)
    elif data.manuscript and data.manuscript.path:
        # The reviewed text without its revision marks: the Word file is for reading, not the project's history.
        reviewed = (base / data.manuscript.path).resolve()
        markdown, edits = viz.revision_marks(reviewed.read_text(encoding="utf-8"), reviewed.suffix.lower(), reviewed)[0], []
    else:
        raise ValueError("The inventory names no reviewed text and no corrected version was given.")
    path = base / data.manuscript.path if data.manuscript and data.manuscript.path else base
    builder = Builder(viz.image_folders(base, path))
    builder.render(markdown)
    changes = {c.id: c for c in data.changes}
    counts = {"problems": 0, "problems_on_their_section": 0, "questions": 0, "reasons": 0}
    if decisions_only:
        counts |= {"decisions": 0, "accepted_edits": len(dict.fromkeys(accepted))}

    by_edit: dict[int, list] = {}
    for run, kind, number, _ in builder.runs:
        if number is not None:
            by_edit.setdefault(number, []).append(run)
    asked = {edit["n"]: asked_for.get(tuple(edit["ids"]), edit.get("comments", [])) for edit in edits}
    for edit in edits:
        runs = by_edit.get(edit["n"])
        if not runs:
            continue
        if questions := asked.get(edit["n"], []):
            # One comment per edit, however many questions it asks, so that a passage is not buried under them.
            if len(questions) == 1:
                builder.comment(runs, f"Question for the authors: {questions[0]} ({', '.join(edit['ids'])})", "Questions for the authors")
            else:
                comment = builder.doc.add_comment(runs, text=f"{len(questions)} questions for the authors about this edit "
                                                  f"({', '.join(edit['ids'])}):", author="Questions for the authors", initials="Q")
                for number, question in enumerate(questions, 1):
                    comment.add_paragraph(f"{number}. {question}")
            counts["questions"] += len(questions)
        if reasons:
            lines = []
            for cid in edit["ids"]:
                change = changes.get(cid)
                if not change or (decisions_only and change.status == "applied"):
                    continue
                why = change.why or " ".join(f"{f.summary} ({f.id}). {f.detail}".strip() for f in data.findings if f.id in change.finding_ids)
                lines.append(f"{'Applied' if change.status == 'applied' else 'Proposed'}: {change.summary.rstrip('.')}."
                             + (f" Why: {why}" if why else "") + f" ({cid})")
            if lines:
                builder.comment(runs, " ".join(lines), stage_of(edit["ids"][0]))
                counts["reasons"] += 1

    # Questions on proposals that are not in this version (each waits for its answer): on the passage they concern.
    waiting: dict[str, list[dict]] = {}
    for item in (revision or {}).get("comments", []):
        if not item.get("placed"):
            waiting.setdefault(", ".join(item["ids"]), []).append(item)
    for ids, items in waiting.items():
        runs = builder.locate(items[0]["passage"]) if items[0].get("passage") else None
        anchor = runs or ([a] if (a := builder.heading_for(items[0]["section"])) is not None else None)
        if anchor is None:
            continue
        # The proposal by its section and summary; its ID follows in parentheses.
        first = changes.get(items[0]["ids"][0])
        named = f"the proposal \u201c{first.summary.rstrip('.')}\u201d in {first.section}" if first else "a proposal"
        comment = builder.doc.add_comment(anchor, text=f"{plural(len(items), 'question')} for the authors about {named}, left out "
                                          f"of this version until they answer ({ids}):", author="Questions for the authors", initials="Q")
        for number, item in enumerate(items, 1):
            comment.add_paragraph(f"{number}. {item['question']}")
        counts["questions"] += len(items)

    answered: set[str] = set()  # problems the decisions document leaves to the workbook: an edit or a decision answers them
    if decisions_only:
        in_text = {cid for piece in (revision or {}).get("pieces", []) if "ids" in piece for cid in piece["ids"]}
        answered = {fid for cid in in_text if cid in changes for fid in changes[cid].finding_ids}
        findings = {f.id: f for f in data.findings}
        for decision in viz.decision_list(data):
            covered = [findings[fid] for fid in decision["finding_ids"] if fid in findings]
            answered |= {f.id for f in covered}
            anchor = next((runs for f in covered if f.quote and (runs := builder.locate(f.quote))), None)
            if anchor is None and (heading := builder.heading_for(covered[0].section if covered else "")) is not None:
                anchor = [heading]
            if anchor is None:
                continue
            builder.comment(anchor, f"A decision only you can make: {decision['summary'].rstrip('.')}. "
                            f"({', '.join(dict.fromkeys([decision['id'], *decision['finding_ids']]))})", "Decisions for the authors")
            counts["decisions"] += 1

    if problems:
        elsewhere: dict[int, tuple[object, list[str]]] = {}  # one comment per heading for problems not in the text
        for finding in data.findings:
            if finding.assessment == "dismissed":
                continue
            if decisions_only and (finding.id in answered or finding.resolution == "addressed"
                                   or finding.severity not in NEEDS_ATTENTION):
                continue
            # What the problem is first; its severity and when to fix it next; its ID last, in parentheses.
            text = (f"{finding.summary.rstrip('.')} ({SEVERITY[finding.severity]}, {WHEN[finding.stage]}"
                    f"{'; a candidate to verify' if finding.assessment == 'candidate' else ''})."
                    + (f" {finding.detail}" if finding.detail else "") + (f" Fix: {finding.action}" if finding.action else "")
                    + f" Found by: {', '.join(skill_name(s) for s in finding.skills)}. ({finding.id})")
            runs = builder.locate(finding.quote) if finding.quote else None
            if runs is None:
                anchor = builder.heading_for(finding.section)
                if anchor is not None:
                    elsewhere.setdefault(id(anchor), (anchor, []))[1].append(f"About {finding.section}: {text}")
                    counts["problems_on_their_section"] += 1
                continue
            builder.comment(runs, text, skill_name(finding.skills[0]))
            counts["problems"] += 1
        for anchor, texts in elsewhere.values():
            comment = builder.doc.add_comment([anchor], text=f"{plural(len(texts), 'problem')} about this part whose passage is not "
                                              "quoted in this text:", author="Workflow", initials="W")
            for line in texts:
                comment.add_paragraph(line)
            counts["problems"] += len(texts)

    for run, source in builder.missing_images:
        name = Path(source.split("#")[0].split("?")[0]).name or source
        builder.comment([run], f"Image not embedded: {name} was not found beside the manuscript, or is not an image Word "
                        "can show.", "Workflow")
    counts["images_not_embedded"] = len(builder.missing_images)

    author_of = {edit["n"]: stage_of(edit["ids"][0]) for edit in edits}
    counts["tracked_changes"] = builder.track(author_of, datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"))
    builder.doc.core_properties.title = data.title
    builder.doc.core_properties.author = "manuscript-improve"
    return builder, counts


@app.command()
def main(inventory: Annotated[Path, typer.Argument(exists=True, dir_okay=False)],
         output: Annotated[Path, typer.Option("--output", help="The .docx to write")],
         revision: Annotated[Path | None, typer.Option(help="revise.py's output folder: the corrected version as tracked changes",
                                                       exists=True, file_okay=False)] = None,
         problems: Annotated[bool, typer.Option(help="A comment for every problem, on the passage it quotes")] = True,
         reasons: Annotated[bool, typer.Option(help="A comment with each edit's reasoning")] = True,
         detail: Annotated[str, typer.Option(help=(
             "decisions (the default): applied edits already in the text, comments only for questions, proposals, "
             "decisions and open major problems; everything: every edit tracked and every problem a comment, for "
             "authors who approve each change (edits: propose)"))] = "decisions") -> None:
    if output.exists():
        raise typer.BadParameter(f"{output.name} exists; write the next attempt as {viz.next_attempt(output, folder=False).name}.")
    if output.suffix.lower() != ".docx":
        raise typer.BadParameter("The output is a .docx file.")
    try:
        data = viz.Inventory.model_validate_json(inventory.read_text(encoding="utf-8"))
        corrected = json.loads((revision / "revision.json").read_text(encoding="utf-8")) if revision else None
        if corrected and corrected.get("inventory_sha256") != hashlib.sha256(inventory.read_bytes()).hexdigest():
            raise ValueError("The corrected version was built from another inventory; rebuild it with revise.py from this one.")
        builder, counts = build(data, inventory.resolve().parent, corrected, problems, reasons, detail)
    except (ValueError, OSError) as exc:
        raise typer.BadParameter(str(exc)) from exc
    output.parent.mkdir(parents=True, exist_ok=True)
    builder.doc.save(str(output))
    typer.echo(json.dumps({"output": str(output.resolve()), **counts}, indent=2))


if __name__ == "__main__":
    for stream in (sys.stdout, sys.stderr):  # UTF-8 on every platform, a Windows console included
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    app()
