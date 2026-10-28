"""The annotated Word document: the corrected version as tracked changes signed by each stage, the review as
comments, and no note to the authors in the text."""
from __future__ import annotations

import importlib.util
import json
import re
import sys
import zipfile
from pathlib import Path

import openpyxl
from typer.testing import CliRunner

ROOT = Path(__file__).resolve().parents[1]


def load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


viz = load("dashboard", ROOT / "skills/manuscript-dashboard/scripts/dashboard.py")
revise = load("revise", ROOT / "skills/manuscript-improve/scripts/revise.py")
export = load("export_docx", ROOT / "skills/manuscript-dashboard/scripts/export_docx.py")

ORIGINAL = """# A tool

## 1. Introduction

Genomes are cheap. This tool plays a crucial role in annotation. We use HuggingFace for data.

| Engine | Time |
|---|---|
| PLINK2 | 5.7 s |

## 2. Methods

The pipeline joins by rsID. It is fast. See [the code](https://example.org/code).
"""


def evidence() -> list[dict]:
    return [{"path": "review.md", "quote": "The review."}]


def make_run(tmp_path: Path) -> Path:
    (tmp_path / "review.md").write_text("The review.")
    (tmp_path / "original.md").write_text(ORIGINAL)
    working = ORIGINAL.replace("HuggingFace", "Hugging Face")
    (tmp_path / "working.md").write_text(working)
    inventory = {
        "schema_version": 2, "title": "Review", "run_id": "run", "scope": "Whole paper.",
        "coverage": [{"skill": skill, "status": "completed", "note": "All.", "evidence": evidence()}
                     for skill in ("manuscript-improve", "manuscript-integrity-check", "manuscript-journal-select")],
        "findings": [
            {"id": "IC-1", "category": "Text versus code", "section": "2. Methods", "summary": "The join key contradicts the code",
             "severity": "major", "assessment": "confirmed", "resolution": "open", "stage": "now", "detail": "The code joins by position.",
             "action": "Say position and genotype.", "skills": ["manuscript-integrity-check"], "evidence": evidence(),
             "quote": "The pipeline joins by rsID."},
            {"id": "JS-1", "category": "Journal fit", "section": "Cover letter", "summary": "No cover letter yet",
             "severity": "minor", "assessment": "confirmed", "resolution": "open", "skills": ["manuscript-journal-select"],
             "evidence": evidence()},
        ],
        "manuscript": {"title": "A tool", "path": "original.md", "sections": ["1. Introduction", "2. Methods"]},
        "changes": [
            {"id": "prose-edit.C1", "kind": "Applied edit", "section": "1. Introduction", "summary": "Spell the company's name",
             "status": "applied", "before": "We use HuggingFace for data.", "after": "We use Hugging Face for data.",
             "evidence": evidence()},
            {"id": "IC-C1", "kind": "Proposed edit", "section": "2. Methods", "summary": "Name the real join key", "status": "proposed",
             "finding_ids": ["IC-1"], "before": "The pipeline joins by rsID.",
             "after": "The pipeline joins by position and genotype [AUTHORS: which genome build?].", "evidence": evidence(),
             "why": "The code joins on position and genotype; rsIDs are not used."},
            {"id": "ST-C1", "kind": "Story edit", "section": "1. Introduction", "summary": "Say what the tool does",
             "status": "proposed", "before": "This tool plays a crucial role in annotation.", "after": "This tool annotates variants.",
             "evidence": evidence()},
            {"id": "JA-C8", "kind": "Proposed edit", "section": "2. Methods", "summary": "Declare the AI use", "status": "proposed",
             "before": "It is fast.", "after": "It is fast. We used [AUTHORS: tool and version] to draft it.", "evidence": evidence()},
        ],
    }
    (tmp_path / "findings.json").write_text(json.dumps(inventory))
    result = CliRunner().invoke(revise.app, [str(tmp_path / "original.md"), str(tmp_path / "working.md"), str(tmp_path / "findings.json"),
                                             "--output-dir", str(tmp_path / "revision")])
    assert result.exit_code == 0, result.output
    return tmp_path


def parts(path: Path) -> tuple[str, str]:
    with zipfile.ZipFile(path) as docx:
        return docx.read("word/document.xml").decode(), docx.read("word/comments.xml").decode()


def test_the_whole_review_is_tracked_changes_with_every_problem_as_a_comment(tmp_path: Path) -> None:
    run = make_run(tmp_path)
    out = run / "review.docx"
    result = CliRunner().invoke(export.app, [str(run / "findings.json"), "--revision", str(run / "revision"), "--output", str(out),
                                             "--detail", "everything"])
    assert result.exit_code == 0, result.output
    counts = json.loads(result.output)
    assert counts["questions"] == 3 and counts["problems"] == 2 and counts["problems_on_their_section"] == 1 and counts["reasons"] == 3
    document, comments = parts(out)
    text = re.sub(r"<[^>]+>", "", document)
    assert "AUTHORS" not in text and "position and genotype" in text
    authors = set(re.findall(r'<w:(?:ins|del) [^>]*w:author="([^"]+)"', document))
    assert authors == {"Prose edit", "Integrity check", "Story"}
    assert re.search(r'<w:del [^>]*><w:r>.*?<w:delText[^>]*>rsID', document)
    # Each comment says what it is about first; the ID follows in parentheses.
    assert "Question for the authors: which genome build? (IC-C1)" in comments
    assert "The join key contradicts the code (Major, fix now)." in comments and "(IC-1)" in comments
    assert ("2 questions for the authors about the proposal \u201cDeclare the AI use\u201d in 2. Methods, left out of this "
            "version until they answer (JA-C8):") in comments
    assert "Proposed: Name the real join key. Why:" in comments
    assert "Why: The code joins on position and genotype; rsIDs are not used." in comments
    assert "1 problem about this part whose passage is not quoted in this text:" in comments and "About Cover letter" in comments
    assert document.count("<w:commentRangeStart") == comments.count("<w:comment ")
    assert "<w:hyperlink" in document and "<w:tbl>" in document
    assert CliRunner().invoke(export.app, [str(run / "findings.json"), "--output", str(out)]).exit_code != 0


def test_pandoc_superscripts_become_superscript_runs_and_quotes_still_anchor() -> None:
    builder = export.Builder([])
    builder.render("**Authors:** Kulaga Anton^1,2,†^, Usanov Nikolay^3^ and water, H~2~O.\n")
    runs = {text: run for run, _, _, text in builder.runs}
    assert runs["1,2,†"].font.superscript and runs["3"].font.superscript and runs["2"].font.subscript
    assert not any(run.font.superscript or run.font.subscript for text, run in runs.items() if text not in {"1,2,†", "3", "2"})
    assert "^" not in "".join(runs) and "~" not in "".join(runs)
    # A quote copied from the Markdown, markers included, still finds its words in the document.
    assert builder.locate("Kulaga Anton^1,2,†^, Usanov Nikolay^3^") is not None


def test_body_paragraphs_are_justified_and_headings_tables_code_and_broken_lines_are_not() -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    builder = export.Builder([])
    builder.render("# Title\n\n## 1. Introduction\n\nBody text that wraps.\n\n- A list item.\n\n> A quote.\n\n"
                   "Kulaga Anton\\\nUniversity of Rostock\n\n```\ncode\n```\n\n| A | B |\n|---|---|\n| 1 | 2 |\n")
    alignment = {p.text: p.paragraph_format.alignment for p in builder.doc.paragraphs}
    assert all(alignment[text] == WD_ALIGN_PARAGRAPH.JUSTIFY for text in ("Body text that wraps.", "A list item.", "A quote."))
    assert alignment["Title"] is None and alignment["1. Introduction"] is None and alignment["code"] is None
    assert alignment["Kulaga Anton\nUniversity of Rostock"] is None
    assert builder.doc.tables[0].cell(1, 0).paragraphs[0].paragraph_format.alignment is None


def test_a_pandoc_attribute_block_after_an_image_is_its_alt_text_and_never_printed(tmp_path: Path) -> None:
    from PIL import Image

    Image.new("RGB", (8, 8), "white").save(tmp_path / "figure-1.png")
    markdown = ('![](figure-1.png){width=100% alt="Two screenshots, {not} nested"}\n\n'
                '![Bracketed alt](figure-1.png) and text after it.\n')
    html_out = viz.markdown_parser().render(markdown)
    assert "{width" not in html_out and 'alt="Two screenshots, {not} nested"' in html_out
    assert "width=100%" not in html_out and 'alt="Bracketed alt"' in html_out and "and text after it." in html_out
    builder = export.Builder([tmp_path])
    builder.render('![](figure-1.png){width=100% alt="Two screenshots of the interface"}\n')
    out = tmp_path / "figure.docx"
    builder.doc.save(out)
    with zipfile.ZipFile(out) as docx:
        document = docx.read("word/document.xml").decode()
    assert "width=100%" not in document and "{" not in re.sub(r"<[^>]+>", "", document)
    assert 'descr="Two screenshots of the interface"' in document and not builder.missing_images


def test_without_a_corrected_version_the_reviewed_text_carries_the_problems_only(tmp_path: Path) -> None:
    run = make_run(tmp_path)
    out = run / "reviewed.docx"
    result = CliRunner().invoke(export.app, [str(run / "findings.json"), "--output", str(out), "--no-reasons"])
    assert result.exit_code == 0, result.output
    document, comments = parts(out)
    assert "<w:ins " not in document and "<w:del " not in document and "HuggingFace" in re.sub(r"<[^>]+>", "", document)
    assert "(Major, fix now)" in comments and "(IC-1)" in comments and "Question for the authors" not in comments


def test_the_reviewed_text_goes_to_word_without_its_revision_marks(tmp_path: Path) -> None:
    run = make_run(tmp_path)
    (run / "original.md").write_text(ORIGINAL.replace("The pipeline joins", "<!-- @rev C1 rounds=1 by=claude:2026-10-10 -->\nThe pipeline joins")
                                     .replace("(https://example.org/code).\n", "(https://example.org/code).\n<!-- @/rev C1 -->\n"))
    out = run / "reviewed.docx"
    result = CliRunner().invoke(export.app, [str(run / "findings.json"), "--output", str(out)])
    assert result.exit_code == 0, result.output
    document, comments = parts(out)
    text = re.sub(r"<[^>]+>", "", document)
    assert "@rev" not in text and "rounds=" not in text and "The pipeline joins by rsID." in text
    assert "(IC-1)" in comments  # the problem's quote still finds its passage


def test_by_default_the_document_holds_only_what_needs_the_authors(tmp_path: Path) -> None:
    run = make_run(tmp_path)
    out = run / "decisions.docx"
    result = CliRunner().invoke(export.app, [str(run / "findings.json"), "--revision", str(run / "revision"), "--output", str(out)])
    assert result.exit_code == 0, result.output
    counts = json.loads(result.output)
    # The applied spelling fix asks nothing: it is in the text, untracked, with no comment.
    assert counts["accepted_edits"] == 1 and counts["reasons"] == 2 and counts["questions"] == 3
    # IC-1 is answered by the proposal that names the real join key, and JS-1 is minor: no problem comments.
    assert counts["problems"] == 0 and counts["decisions"] == 0
    document, comments = parts(out)
    text = re.sub(r"<[^>]+>", "", document)
    assert "Hugging Face" in text and "HuggingFace" not in text
    authors = set(re.findall(r'<w:(?:ins|del) [^>]*w:author="([^"]+)"', document))
    assert authors == {"Integrity check", "Story"}  # the proposals only
    assert "Applied:" not in comments and "Spell the company" not in comments
    assert "Proposed: Name the real join key." in comments and "Question for the authors: which genome build? (IC-C1)" in comments
    assert "The join key contradicts the code" not in comments and "No cover letter yet" not in comments
    assert document.count("<w:commentRangeStart") == comments.count("<w:comment ")


def test_a_decision_and_an_unanswered_major_problem_are_comments_and_minor_ones_are_not(tmp_path: Path) -> None:
    run = make_run(tmp_path)
    data = json.loads((run / "findings.json").read_text())
    data["findings"] += [
        {"id": "PR-1", "category": "Validation", "section": "2. Methods", "summary": "No known-answer test of the join",
         "severity": "major", "assessment": "confirmed", "resolution": "open", "stage": "now", "skills": ["manuscript-peer-review"],
         "evidence": evidence(), "quote": "It is fast.", "decision": True},
        {"id": "PR-2", "category": "Benchmark", "section": "2. Methods", "summary": "Speed shown on one genome only",
         "severity": "major", "assessment": "confirmed", "resolution": "open", "stage": "now", "skills": ["manuscript-peer-review"],
         "evidence": evidence(), "quote": "It is fast."},
        {"id": "DS-1", "category": "Writing", "section": "1. Introduction", "summary": "Cheap is vague",
         "severity": "minor", "assessment": "confirmed", "resolution": "open", "stage": "writing", "skills": ["manuscript-deslop"],
         "evidence": evidence(), "quote": "Genomes are cheap."},
    ]
    data["decisions"] = [{"id": "D1", "summary": "Run a known-answer test or narrow the claim", "finding_ids": ["PR-1"]}]
    data["coverage"] += [{"skill": skill, "status": "completed", "note": "All.", "evidence": evidence()}
                         for skill in ("manuscript-peer-review", "manuscript-deslop")]
    revision = json.loads((run / "revision" / "revision.json").read_text())
    builder, counts = export.build(viz.Inventory.model_validate(data), run, revision, True, True, "decisions")
    out = run / "with-decisions.docx"
    builder.doc.save(str(out))
    _, comments = parts(out)
    assert counts["decisions"] == 1 and counts["problems"] == 1
    assert "A decision only you can make: Run a known-answer test or narrow the claim. (D1, PR-1)" in comments
    assert "Speed shown on one genome only (Major, fix now)." in comments  # open, major, nothing answers it
    assert "No known-answer test of the join" not in comments.split("(D1, PR-1)")[1]  # the decision covers it
    assert "Cheap is vague" not in comments  # minor: in the workbook, not the document


def test_the_changes_workbook_lists_applied_edits_and_what_needs_the_authors(tmp_path: Path) -> None:
    changes = load("export_changes", ROOT / "skills/manuscript-dashboard/scripts/export_changes.py")
    run = make_run(tmp_path)
    out = run / "manuscript-changes.xlsx"
    args = [str(run / "findings.json"), "--revision", str(run / "revision"), "--output", str(out)]
    result = CliRunner().invoke(changes.app, args)
    assert result.exit_code == 0, result.output
    counts = json.loads(result.output)
    assert counts["applied"] == 1 and counts["restored"] == 0
    book = openpyxl.load_workbook(out)
    assert book.sheetnames == ["Applied", "For you", "Open problems", "Held"]
    rows = {name: [r for r in book[name].iter_rows(min_row=2, values_only=True)] for name in book.sheetnames}
    [applied] = rows["Applied"]
    assert applied[0] == "prose-edit.C1" and applied[1] == "Prose edit" and applied[5] == "We use HuggingFace for data."
    assert applied[6] == "We use Hugging Face for data." and applied[4] == "applied"
    kinds = {r[1]: r[0] for r in rows["For you"]}
    assert kinds == {"IC-C1": "proposal in this version", "ST-C1": "proposal in this version",
                     "JA-C8": "proposal waiting for your answer"}
    assert any("Question: which genome build?" in (r[6] or "") for r in rows["For you"])
    # Every problem no applied edit answers, whatever its severity, and the proposal that would answer it.
    assert {r[0]: r[8] for r in rows["Open problems"]} == {"IC-1": "IC-C1", "JS-1": None}
    assert [r[0] for r in rows["Held"]] == ["JA-C8"]  # left out of this version until the authors answer
    assert book["Applied"].freeze_panes == "A2" and book["Applied"].auto_filter.ref
    refused = CliRunner().invoke(changes.app, args)
    assert refused.exit_code != 0 and "manuscript-changes.2.xlsx" in refused.output


def test_the_changes_workbook_adds_the_co_authors_comments_only_when_the_run_tracked_them(tmp_path: Path) -> None:
    changes = load("export_changes", ROOT / "skills/manuscript-dashboard/scripts/export_changes.py")
    run = make_run(tmp_path)
    data = json.loads((run / "findings.json").read_text())
    revision = json.loads((run / "revision" / "revision.json").read_text())
    book, counts = changes.build(viz.Inventory.model_validate(data), revision)
    assert book.sheetnames == ["Applied", "For you", "Open problems", "Held"] and counts["coauthors"] == 0
    data["coauthors"] = {"round": "0.5", "next_round": "0.6", "source": "comments.yaml", "items": [
        {"id": "C1", "kind": "comment", "author": "Robi Tacutu", "where": "2. Methods", "comment": "Which key do we join on?",
         "status": "partly", "before": "The pipeline joins by rsID.", "after": "The pipeline joins by position. It is fast.",
         "now": "The pipeline joins by position.", "why": "The key is named.", "next": "Say which genome build.",
         "finding_ids": ["IC-1"]},
        {"id": "S1", "kind": "suggestion", "author": "Olga Ivanova", "where": "1. Introduction",
         "edits": [{"deleted": "cheap", "inserted": "inexpensive"}], "status": "declined", "why": "Cheap is the plainer word."}]}
    book, counts = changes.build(viz.Inventory.model_validate(data), revision)
    assert book.sheetnames == ["Applied", "For you", "Open problems", "Held", "Co-authors' comments"] and counts["coauthors"] == 2
    out = run / "with-coauthors.xlsx"
    book.save(str(out))
    sheet = openpyxl.load_workbook(out)["Co-authors' comments"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("Id", "Author", "Kind", "Section", "Comment", "Status", "Was", "Now", "Why", "What could be done")
    assert rows[1] == ("C1", "Robi Tacutu", "Comment", "2. Methods", "Which key do we join on?", "Partly addressed",
                       "The pipeline joins by rsID.", "The pipeline joins by position.", "The key is named.", "Say which genome build.")
    assert rows[2][:6] == ("S1", "Olga Ivanova", "Suggested edit", "1. Introduction", "Delete: “cheap” Insert: “inexpensive”", "Declined")
    assert rows[2][8:] == ("Cheap is the plainer word.", None)
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref
